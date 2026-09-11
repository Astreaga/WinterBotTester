import sys
import os
import re
import shutil
import copy
import csv
import time
import json
import urllib.request
import urllib.error
import subprocess
import threading
from pathlib import Path
from datetime import datetime
from collections import defaultdict

import pandas as pd
from docx import Document
from playwright.sync_api import sync_playwright, TimeoutError as PlaywrightTimeoutError

from PySide6.QtCore import QThread, Signal, QSettings
from PySide6.QtWidgets import (
    QApplication, QMainWindow, QWidget, QVBoxLayout, QHBoxLayout,
    QLabel, QLineEdit, QPushButton, QFileDialog, QCheckBox,
    QTextEdit, QTableWidget, QTableWidgetItem, QHeaderView,
    QMessageBox, QProgressBar, QFrame, QSpinBox, QComboBox
)


TEXT_BOX = "#chat-input"
SEND_BUTTON = "#send-message-button"
WORD_BUTTON = 'button[aria-label="Download as a Word Document"]'
MICROSOFT_BUTTON_TEXT = "Continue with Microsoft"

APP_NAME = "Winter Bot Tester"
APP_VERSION = "1.3 (encoding-safe rating)"

# Allowed feedback reasons, taken from the Expedient feedback panel.
POSITIVE_REASONS = [
    "Accurate information", "Followed instructions perfectly",
    "Showcased creativity", "Positive attitude",
    "Attention to detail", "Thorough explanation", "Other"
]
NEGATIVE_REASONS = [
    "Don't like the style", "Too verbose", "Not helpful",
    "Not factually correct", "Didn't fully follow instructions",
    "Refused when it shouldn't have", "Being lazy", "Other"
]

# Feedback panel selectors. The thumbs icons MUST be confirmed by inspecting the
# page (right-click a thumb -> Inspect). These are best-guess defaults; everything
# else is targeted by visible text/placeholder, which is far more stable.
THUMBS_UP_SELECTOR = 'button[aria-label="Good Response"]'
THUMBS_DOWN_SELECTOR = 'button[aria-label="Bad Response"]'
RATING_PANEL_TEXT = "How would you rate this response?"
DETAILS_PLACEHOLDER = "Feel free to add specific details"
TAG_PLACEHOLDER = "Add a tag"
SAVE_BUTTON_TEXT = "Save"

OPENAI_URL = "https://api.openai.com/v1/chat/completions"


def app_support_dir() -> Path:
    """A stable, writable folder for app data on macOS.

    Inside a packaged .app the working directory is read-only and
    unpredictable, so the Chromium login profile must live here instead
    of next to the executable.
    """
    base = Path.home() / "Library" / "Application Support" / APP_NAME
    base.mkdir(parents=True, exist_ok=True)
    return base


def ensure_chromium(log):
    """Make sure Playwright's Chromium is installed; download it on first run.

    On the machine you build on, Chromium is already present from
    `playwright install chromium`, so this is a no-op. On a fresh Mac that
    received the .app, it downloads Chromium once into the user cache.
    """
    try:
        with sync_playwright() as p:
            exe = p.chromium.executable_path
        if exe and Path(exe).exists():
            return
    except Exception as exc:
        log(f"Browser check failed, will attempt install: {exc}")

    log("Setting up the browser engine (first run only).")
    log("This downloads ~150 MB and can take a few minutes...")

    if getattr(sys, "frozen", False):
        # Packaged app: use Playwright's bundled Node driver.
        from playwright._impl._driver import compute_driver_executable
        try:
            from playwright._impl._driver import get_driver_env
            env = get_driver_env()
        except Exception:
            env = os.environ.copy()

        driver = compute_driver_executable()
        if isinstance(driver, (list, tuple)):
            cmd = list(driver) + ["install", "chromium"]
        else:
            cmd = [driver, "install", "chromium"]
    else:
        # Running from source: plain python -m playwright works.
        cmd = [sys.executable, "-m", "playwright", "install", "chromium"]
        env = os.environ.copy()

    proc = subprocess.Popen(
        cmd, env=env,
        stdout=subprocess.PIPE, stderr=subprocess.STDOUT, text=True
    )
    for line in proc.stdout:
        line = line.rstrip()
        if line:
            log(line)
    proc.wait()

    if proc.returncode != 0:
        raise RuntimeError(
            f"Could not set up the browser engine (exit code {proc.returncode})."
        )

    log("Browser engine ready.")


def safe_name(value: str) -> str:
    value = value.strip()
    value = re.sub(r'[\\/:*?"<>|]+', "-", value)
    value = re.sub(r"\s+", " ", value)
    return value[:120].strip()


def load_questions(excel_path: str):
    df = pd.read_excel(excel_path)
    df.columns = [str(c).strip() for c in df.columns]

    required = {"Supplier", "Question", "Repetitions"}
    missing = required - set(df.columns)

    if missing:
        raise ValueError(f"Missing columns: {', '.join(missing)}")

    jobs = []
    supplier_question_counter = defaultdict(int)

    has_approved = "Approved Answer" in df.columns
    has_expected = "Expected Suppliers" in df.columns

    for _, row in df.iterrows():
        supplier = str(row["Supplier"]).strip()
        question = str(row["Question"]).strip()
        repetitions = int(row["Repetitions"])

        if not supplier or supplier.lower() == "nan":
            raise ValueError("Supplier cannot be empty.")

        if not question or question.lower() == "nan":
            raise ValueError("Question cannot be empty.")

        approved = ""
        if has_approved:
            approved = str(row["Approved Answer"]).strip()
            if approved.lower() == "nan":
                approved = ""

        expected_suppliers = ""
        if has_expected:
            expected_suppliers = str(row["Expected Suppliers"]).strip()
            if expected_suppliers.lower() == "nan":
                expected_suppliers = ""

        supplier_question_counter[supplier] += 1

        jobs.append({
            "supplier": supplier,
            "question_index": supplier_question_counter[supplier],
            "question": question,
            "repetitions": repetitions,
            "approved": approved,
            "expected_suppliers": expected_suppliers
        })

    return jobs


def group_jobs_by_supplier(jobs):
    grouped = defaultdict(list)

    for job in jobs:
        grouped[job["supplier"]].append(job)

    return grouped


def build_target_path(output_root, supplier, question_index, round_number, extension=".docx"):
    date_str = datetime.now().strftime("%Y-%m-%d")
    supplier_clean = safe_name(supplier)

    folder = Path(output_root) / supplier_clean / f"Q{question_index:02d}"
    folder.mkdir(parents=True, exist_ok=True)

    filename = f"{supplier_clean}_Q{question_index:02d}_Round{round_number:02d}_{date_str}{extension}"
    return folder / filename


def build_diagnostic_paths(output_root, supplier, question_index, round_number):
    """Paths for a screenshot + HTML snapshot saved when a run fails."""
    ts = datetime.now().strftime("%Y-%m-%d_%H%M%S")
    supplier_clean = safe_name(supplier)
    folder = Path(output_root) / supplier_clean / "_diagnostics"
    folder.mkdir(parents=True, exist_ok=True)
    stem = f"{supplier_clean}_Q{question_index:02d}_Round{round_number:02d}_FAIL_{ts}"
    return folder / f"{stem}.png", folder / f"{stem}.html"


def existing_round_file(output_root, supplier, question_index, round_number):
    """Return today's already-downloaded file for this run, if any (for resume)."""
    date_str = datetime.now().strftime("%Y-%m-%d")
    supplier_clean = safe_name(supplier)
    folder = Path(output_root) / supplier_clean / f"Q{question_index:02d}"

    if not folder.exists():
        return None

    pattern = f"{supplier_clean}_Q{question_index:02d}_Round{round_number:02d}_{date_str}.*"
    matches = sorted(folder.glob(pattern))
    return matches[0] if matches else None


def docx_has_text(path):
    """True if the Word file contains any real text (paragraphs or tables)."""
    try:
        doc = Document(path)
    except Exception:
        return False

    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for table_row in table.rows:
            for cell in table_row.cells:
                parts.append(cell.text)

    return any(s and s.strip() for s in parts)


def find_round_file_any_date(output_root, supplier, question_index, round_number):
    """Newest downloaded Word file for this round, regardless of date."""
    supplier_clean = safe_name(supplier)
    folder = Path(output_root) / supplier_clean / f"Q{question_index:02d}"

    if not folder.exists():
        return None

    matches = [
        m for m in folder.glob(
            f"{supplier_clean}_Q{question_index:02d}_Round{round_number:02d}_*"
        )
        if m.suffix.lower() == ".docx"
    ]
    matches.sort(key=lambda p: p.stat().st_mtime, reverse=True)
    return matches[0] if matches else None


def append_result_row(output_root, supplier, question_index, question, round_number, status, detail):
    """Append one line to a per-day results CSV in the output folder."""
    date_str = datetime.now().strftime("%Y-%m-%d")
    csv_path = Path(output_root) / f"results_{date_str}.csv"
    fieldnames = ["timestamp", "supplier", "question_index", "question",
                  "round", "status", "detail"]
    write_header = not csv_path.exists()

    with open(csv_path, "a", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=fieldnames)
        if write_header:
            writer.writeheader()
        writer.writerow({
            "timestamp": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
            "supplier": supplier,
            "question_index": question_index,
            "question": question,
            "round": round_number,
            "status": status,
            "detail": str(detail),
        })


def copy_docx_body(source_path: Path, target_doc: Document):
    source_doc = Document(source_path)

    target_body = target_doc.element.body
    target_sectPr = target_body.sectPr

    for element in source_doc.element.body:
        if element.tag.endswith("sectPr"):
            continue

        copied_element = copy.deepcopy(element)

        if target_sectPr is not None:
            target_body.insert(target_body.index(target_sectPr), copied_element)
        else:
            target_body.append(copied_element)


def create_supplier_combined_report(output_root, supplier, supplier_results):
    date_str = datetime.now().strftime("%Y-%m-%d")
    supplier_clean = safe_name(supplier)
    supplier_folder = Path(output_root) / supplier_clean
    supplier_folder.mkdir(parents=True, exist_ok=True)

    report_path = supplier_folder / f"{supplier_clean}_Combined_Report_{date_str}.docx"

    document = Document()
    document.add_heading(f"{supplier} - Combined Bot Test Report", level=1)
    document.add_paragraph(f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")

    separator = "=" * 65

    for question_data in supplier_results:
        question_index = question_data["question_index"]
        question_text = question_data["question"]
        rounds = question_data["rounds"]

        document.add_paragraph(separator)
        document.add_heading(f"QUESTION {question_index}", level=1)
        document.add_paragraph(separator)

        prompt_paragraph = document.add_paragraph()
        prompt_paragraph.add_run("PROMPT: ").bold = True
        prompt_paragraph.add_run(question_text)

        for round_number, file_path in rounds:
            document.add_paragraph("")
            round_paragraph = document.add_paragraph()
            round_paragraph.add_run(f"----- ROUND {round_number} -----").bold = True

            try:
                copy_docx_body(Path(file_path), document)
            except Exception as exc:
                document.add_paragraph(f"[ERROR COPYING FILE: {file_path}]")
                document.add_paragraph(str(exc))

    document.save(report_path)
    return report_path


def supplier_has_downloads(output_root, supplier, supplier_jobs):
    """True if at least one round file exists for this supplier."""
    for job in supplier_jobs:
        for round_number in range(1, job["repetitions"] + 1):
            if find_round_file_any_date(output_root, supplier, job["question_index"], round_number):
                return True
    return False


def build_combined_report_from_spec(output_root, supplier, supplier_jobs):
    """Build the combined report from the question spec + whatever files exist on
    disk. Works correctly after partial runs, skips, and retries."""
    date_str = datetime.now().strftime("%Y-%m-%d")
    supplier_clean = safe_name(supplier)
    supplier_folder = Path(output_root) / supplier_clean
    supplier_folder.mkdir(parents=True, exist_ok=True)

    report_path = supplier_folder / f"{supplier_clean}_Combined_Report_{date_str}.docx"

    document = Document()
    document.add_heading(f"{supplier} - Combined Bot Test Report", level=1)
    document.add_paragraph(f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")

    separator = "=" * 65

    for job in supplier_jobs:
        question_index = job["question_index"]
        question_text = job["question"]
        repetitions = job["repetitions"]

        document.add_paragraph(separator)
        document.add_heading(f"QUESTION {question_index}", level=1)
        document.add_paragraph(separator)

        prompt_paragraph = document.add_paragraph()
        prompt_paragraph.add_run("PROMPT: ").bold = True
        prompt_paragraph.add_run(question_text)

        for round_number in range(1, repetitions + 1):
            document.add_paragraph("")
            round_paragraph = document.add_paragraph()
            round_paragraph.add_run(f"----- ROUND {round_number} -----").bold = True

            file_path = find_round_file_any_date(
                output_root, supplier, question_index, round_number
            )

            if file_path:
                try:
                    copy_docx_body(file_path, document)
                except Exception as exc:
                    document.add_paragraph(f"[ERROR COPYING FILE: {file_path}]")
                    document.add_paragraph(str(exc))
            else:
                document.add_paragraph("[No successful download for this round]")

    document.save(report_path)
    return report_path


def docx_extract_text(path):
    """Full visible text of a Word file, in document order, with tables rendered
    as readable rows (cells joined by ' | '). This matters because the bot puts
    its scorecards and recommendation tables in Word tables; if those aren't
    included, an AI reviewer only sees the narration and thinks no
    recommendations were made."""
    try:
        doc = Document(path)
    except Exception:
        return ""

    from docx.oxml.ns import qn
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    parts = []
    for child in doc.element.body.iterchildren():
        if child.tag == qn("w:p"):
            para = Paragraph(child, doc)
            text = para.text.strip()
            if text:
                parts.append(text)
        elif child.tag == qn("w:tbl"):
            table = Table(child, doc)
            rows_out = []
            for row in table.rows:
                cells = [c.text.strip().replace("\n", " ") for c in row.cells]
                if any(cells):
                    rows_out.append(" | ".join(cells))
            if rows_out:
                parts.append("\n".join(rows_out))
    return "\n".join(parts)


_PUNCT_MAP = {
    "\u2014": "-", "\u2013": "-",        # em / en dash
    "\u2018": "'", "\u2019": "'",        # curly single quotes
    "\u201c": '"', "\u201d": '"',        # curly double quotes
    "\u2026": "...",                      # ellipsis
    "\u00a0": " ",                        # non-breaking space
    "\u2022": "*",                        # bullet
}


def encode_safe(text):
    """Make text safe to send over any encoding path (e.g. proxies that use
    latin-1) by replacing typographic characters and dropping anything that
    can't be represented. The bot answers use em-dashes, which broke the request."""
    if not text:
        return text
    for bad, good in _PUNCT_MAP.items():
        text = text.replace(bad, good)
    return text.encode("latin-1", "ignore").decode("latin-1")


def _normalize(text):
    """Lowercase and collapse non-alphanumerics, for forgiving name matching."""
    return re.sub(r"[^a-z0-9]+", " ", (text or "").lower()).strip()


def supplier_present(name, text):
    """True if a supplier name surfaces anywhere in the answer text
    (case- and punctuation-insensitive)."""
    name = (name or "").strip()
    if not name:
        return False
    return _normalize(name) in _normalize(text)


def ai_judge(api_key, model, question, approved, expected_suppliers, bot_answer,
             siblings, target_supplier=""):
    """Verdict based ONLY on whether the right supplier(s) SURFACE in the answer.

    The supplier under test (target_supplier, i.e. the folder the question lives
    in) must appear -> thumbs up; otherwise thumbs down. Writing quality,
    verbosity, structure and instruction-following are NOT judged. Consistency
    across earlier rounds is reported. The OpenAI key is OPTIONAL and used only to
    phrase the feedback sentence; the decision/score are fully deterministic."""

    # Expected list: the supplier under test first, then any extras from the
    # optional "Expected Suppliers" column.
    expected = []
    if target_supplier and target_supplier.strip():
        expected.append(target_supplier.strip())
    for s in (expected_suppliers or "").split(","):
        s = s.strip()
        if s and _normalize(s) not in [_normalize(e) for e in expected]:
            expected.append(s)

    siblings = siblings or []

    found = [s for s in expected if supplier_present(s, bot_answer)]
    missing = [s for s in expected if s not in found]

    primary = expected[0] if expected else ""
    primary_present = supplier_present(primary, bot_answer) if primary else False

    # Cross-round consistency for the primary supplier (this round + earlier ones).
    prior_hits = sum(1 for sib in siblings if primary and supplier_present(primary, sib))
    total_rounds = len(siblings) + 1
    total_hits = prior_hits + (1 if primary_present else 0)
    consistency = total_hits / total_rounds if total_rounds else 1.0

    if not primary:
        return {
            "decision": "down", "score": 1,
            "suppliers_found": found, "suppliers_missing": missing,
            "reasons": ["Other"], "tags": [],
            "feedback": ("No supplier was set to check for this question, so there was "
                         "nothing to test. The supplier under test is normally the folder "
                         "name; you can also fill the Expected Suppliers column."),
            "reasoning": "No expected supplier provided.",
        }

    if primary_present:
        decision = "up"
        score = max(8, min(10, 8 + int(round(2 * consistency))))
        reasons = ["Accurate information"]
        extras = [s for s in found if _normalize(s) != _normalize(primary)]
        extra_note = f" Also surfaced: {', '.join(extras)}." if extras else ""
        miss_note = f" Not surfaced: {', '.join(missing)}." if missing else ""
        feedback = f"{primary} surfaced in this response.{extra_note}{miss_note}".strip()
        reasoning = (f"{primary} appeared in {total_hits} of {total_rounds} rounds so far "
                     f"({int(round(consistency * 100))}% consistent).")
    else:
        decision = "down"
        score = 2
        reasons = ["Not factually correct"]
        feedback = f"{primary} did NOT surface in this response, but it should have."
        reasoning = (f"{primary} appeared in {total_hits} of {total_rounds} rounds so far; "
                     f"it was missing from this one.")

    # OPTIONAL nicety: let OpenAI rephrase the feedback sentence. Never decides
    # anything; skipped entirely if no/!valid key, and any error falls back silently.
    api_key = (api_key or "").strip()
    key_ok = api_key and not (
        len(api_key) > 300 or any(c.isspace() for c in api_key)
        or any(ord(c) > 127 for c in api_key)
    )
    if key_ok:
        try:
            nicer = _ai_feedback_sentence(api_key, model, primary, decision,
                                          found, missing, total_hits, total_rounds)
            if nicer:
                feedback = nicer
        except Exception:
            pass

    return {
        "decision": decision,
        "score": score,
        "suppliers_found": found,
        "suppliers_missing": missing,
        "reasons": reasons,
        "tags": [],
        "feedback": feedback,
        "reasoning": reasoning,
    }


def _ai_feedback_sentence(api_key, model, primary, decision, found, missing,
                          hits, total):
    """Optional: ask OpenAI for ONE short, supplier-focused feedback sentence.
    Returns a string or None. Never affects the verdict."""
    sys_msg = ("You write ONE short, plain feedback sentence for a "
               "supplier-recommendation bot being trained. Comment ONLY on whether the "
               "target supplier surfaced and how consistently; never mention writing "
               "style, length, or formatting. Return ONLY the sentence, no preamble.")
    usr_msg = (f"Target supplier: {primary}\n"
               f"Result: {'surfaced' if decision == 'up' else 'did NOT surface'}\n"
               f"Other expected suppliers found: {', '.join(found) or 'none'}\n"
               f"Expected but missing: {', '.join(missing) or 'none'}\n"
               f"Consistency: appeared in {hits} of {total} rounds.")
    body = {
        "model": model or "gpt-4o-mini",
        "temperature": 0,
        "max_tokens": 80,
        "messages": [
            {"role": "system", "content": encode_safe(sys_msg)},
            {"role": "user", "content": encode_safe(usr_msg)},
        ],
    }
    data = json.dumps(body).encode("utf-8")
    req = urllib.request.Request(
        OPENAI_URL, data=data,
        headers={"Content-Type": "application/json",
                 "Authorization": f"Bearer {api_key}"},
    )
    with urllib.request.urlopen(req, timeout=60) as resp:
        payload = json.loads(resp.read().decode("utf-8"))
    return payload["choices"][0]["message"]["content"].strip()


def append_qa_feedback(path, result):
    """Append an AI QA Feedback section to the downloaded Word file."""
    try:
        document = Document(path)
    except Exception:
        return

    document.add_paragraph("")
    document.add_heading("AI QA Feedback", level=2)

    def line(label, value):
        p = document.add_paragraph()
        p.add_run(f"{label}: ").bold = True
        p.add_run(str(value))

    line("Decision", "Thumbs up" if result["decision"] == "up" else "Thumbs down")
    line("Score", f"{result['score']}/10")
    line("Suppliers found", ", ".join(result["suppliers_found"]) or "none")
    line("Suppliers missing", ", ".join(result["suppliers_missing"]) or "none")
    line("Reasons", ", ".join(result["reasons"]) or "none")
    if result["tags"]:
        line("Tags", ", ".join(result["tags"]))
    line("Feedback", result["feedback"])
    line("Reasoning", result["reasoning"])

    document.save(path)


def _click_score(page, score, log):
    """Try several ways to click the 1-10 score control. Returns True on success."""
    s = str(score)
    strategies = [
        ("button role", lambda: page.get_by_role("button", name=s, exact=True).first),
        ("radio role", lambda: page.get_by_role("radio", name=s, exact=True).first),
        ("aria-label", lambda: page.locator(f'[aria-label="{s}"]').first),
        ("button exact text", lambda: page.locator(f'button:text-is("{s}")').first),
        ("any exact text", lambda: page.get_by_text(s, exact=True).first),
        ("data-score attr", lambda: page.locator(f'[data-score="{s}"]').first),
    ]
    for name, make in strategies:
        try:
            make().click(timeout=2500)
            log(f"Set score {s} (via {name}).")
            return True
        except Exception:
            continue
    return False


def _dump_panel(page, shot_path, log):
    """Save the feedback panel's HTML and a screenshot for diagnosing selectors."""
    base = (shot_path or "").rsplit(".png", 1)[0] or "panel"
    html_path = base + "_panel.html"
    png_path = base + "_panel.png"
    try:
        try:
            html = page.locator('div[role="dialog"]').first.inner_html(timeout=1500)
        except Exception:
            html = page.content()
        with open(html_path, "w", encoding="utf-8") as f:
            f.write(html)
        log(f"Saved panel HTML: {html_path}")
    except Exception:
        pass
    try:
        page.screenshot(path=png_path, full_page=False)
        log(f"Saved panel screenshot: {png_path}")
    except Exception:
        pass


def submit_rating(page, result, log=lambda m: None, shot_path=None):
    """Drive the Expedient feedback panel from the AI result. Logs each step,
    verifies the panel closes after Save (the signal the site accepted it), and
    saves a screenshot afterwards so you can confirm visually."""
    decision = result["decision"]
    thumb = THUMBS_UP_SELECTOR if decision == "up" else THUMBS_DOWN_SELECTOR

    log(f"Clicking thumbs {'up' if decision == 'up' else 'down'}...")
    page.locator(thumb).last.click(timeout=15000)

    page.wait_for_selector(f"text={RATING_PANEL_TEXT}", timeout=15000)
    log("Feedback panel opened.")

    # Score 1-10. The control type varies (button / radio / clickable number),
    # so try several locators before giving up.
    if not _click_score(page, result["score"], log):
        log(f"Could not set score {result['score']} — saving the panel HTML/screenshot "
            f"so the score control can be identified.")
        _dump_panel(page, shot_path, log)

    # Reason buttons.
    for reason in result["reasons"]:
        try:
            page.get_by_role("button", name=reason, exact=True).first.click(timeout=5000)
            log(f"Selected reason: {reason}")
        except Exception as exc:
            log(f"Could not click reason '{reason}': {exc}")

    # Written feedback.
    if result["feedback"]:
        try:
            page.get_by_placeholder(DETAILS_PLACEHOLDER).first.fill(result["feedback"])
            log("Filled feedback details.")
        except Exception as exc:
            log(f"Could not fill feedback: {exc}")

    # Tags (optional).
    for tag in result["tags"]:
        try:
            box = page.get_by_placeholder(TAG_PLACEHOLDER).first
            box.fill(tag)
            box.press("Enter")
        except Exception:
            pass

    # Click Save.
    page.get_by_role("button", name=SAVE_BUTTON_TEXT, exact=True).first.click(timeout=8000)
    log("Clicked Save.")

    # Verify the panel actually closed — that's the real confirmation the site
    # accepted the rating. If it lingers, the save likely didn't register.
    saved = False
    try:
        page.wait_for_selector(f"text={RATING_PANEL_TEXT}", state="hidden", timeout=8000)
        saved = True
    except Exception:
        saved = False

    if saved:
        log(f"Rating SAVED on site: {decision} ({result['score']}/10) — panel closed.")
    else:
        log("WARNING: clicked Save but the feedback panel did not close — "
            "the rating may NOT have been saved. Check the screenshot.")

    # Leave visual proof either way.
    if shot_path:
        try:
            page.screenshot(path=shot_path, full_page=False)
            log(f"Saved rating screenshot: {shot_path}")
        except Exception:
            pass

    return saved


class BotRunner:
    def __init__(self, bot_url, output_root, headless=False, log=None):
        self.bot_url = bot_url
        self.output_root = output_root
        self.headless = headless
        self.log = log or (lambda message: None)
        # Writable, persistent location so the Microsoft login survives
        # between runs even inside a packaged .app.
        self.profile_dir = str(app_support_dir() / "browser_profile")

    def run_one(self, supplier, question_index, question, round_number, rate_callback=None):
        with sync_playwright() as p:
            context = p.chromium.launch_persistent_context(
                user_data_dir=self.profile_dir,
                headless=self.headless,
                accept_downloads=True,
                viewport={"width": 1400, "height": 900}
            )

            page = context.new_page()
            page.set_default_timeout(45_000)

            try:
                page.goto(self.bot_url, wait_until="domcontentloaded")

                try:
                    page.get_by_text(MICROSOFT_BUTTON_TEXT, exact=False).click(timeout=5000)
                    page.wait_for_load_state("domcontentloaded", timeout=30_000)
                except Exception:
                    pass

                page.wait_for_selector(TEXT_BOX, timeout=60_000)
                page.locator(TEXT_BOX).click()
                page.keyboard.insert_text(question)

                page.wait_for_selector(SEND_BUTTON, timeout=30_000)
                page.locator(SEND_BUTTON).click()

                page.wait_for_selector(WORD_BUTTON, timeout=180_000)

                with page.expect_download(timeout=60_000) as download_info:
                    page.locator(WORD_BUTTON).click()

                download = download_info.value
                temp_path = download.path()
                suggested_name = download.suggested_filename or "answer.docx"
                extension = Path(suggested_name).suffix or ".docx"

                target = build_target_path(
                    self.output_root,
                    supplier,
                    question_index,
                    round_number,
                    extension
                )

                shutil.move(temp_path, target)

                # Optional rating step, run while the page is still open so the
                # thumbs/feedback panel can be used if not in dry-run.
                if rate_callback is not None:
                    rate_callback(target, page)

                return target

            except Exception:
                # Save what the page looked like so failures are diagnosable
                # (logged out? bot error? changed layout?) instead of opaque.
                try:
                    png_path, html_path = build_diagnostic_paths(
                        self.output_root, supplier, question_index, round_number
                    )
                    page.screenshot(path=str(png_path), full_page=True)
                    html_path.write_text(page.content(), encoding="utf-8")
                    self.log(f"Saved failure snapshot: {png_path}")
                except Exception:
                    pass
                raise

            finally:
                page.close()
                context.close()


class LoginWorker(QThread):
    """Opens a visible browser on the bot URL and waits for a manual login.

    This is separate from the automated runs: it sets no action timeout and
    waits up to 10 minutes (or until you click "Finish login & save"), so the
    Microsoft sign-in window won't close out from under you. The session is
    stored in the same persistent profile the automated runs use.
    """
    log = Signal(str)
    done = Signal(bool)   # True if the chat box was detected
    failed = Signal(str)

    def __init__(self, bot_url):
        super().__init__()
        self.bot_url = bot_url
        self.finish_requested = False
        self.profile_dir = str(app_support_dir() / "browser_profile")

    def request_finish(self):
        self.finish_requested = True

    def run(self):
        try:
            ensure_chromium(self.log.emit)
            with sync_playwright() as p:
                context = p.chromium.launch_persistent_context(
                    user_data_dir=self.profile_dir,
                    headless=False,            # must be visible to log in
                    accept_downloads=True,
                    viewport={"width": 1400, "height": 900}
                )

                page = context.new_page()
                page.set_default_timeout(0)    # no action timeout during login

                try:
                    self.log.emit("Opening the bot. Sign in with Microsoft in the browser window.")
                    page.goto(self.bot_url, wait_until="domcontentloaded")

                    try:
                        page.get_by_text(MICROSOFT_BUTTON_TEXT, exact=False).click(timeout=5000)
                    except Exception:
                        pass

                    self.log.emit("Waiting for you to finish logging in (up to 10 minutes)...")
                    logged_in = False
                    deadline = time.time() + 600

                    while time.time() < deadline and not self.finish_requested:
                        try:
                            if page.query_selector(TEXT_BOX):
                                logged_in = True
                                break
                        except Exception:
                            pass
                        page.wait_for_timeout(1000)

                    if logged_in:
                        self.log.emit("Login detected. Saving session...")
                    elif self.finish_requested:
                        self.log.emit("Finishing login as requested. Saving session...")
                    else:
                        self.log.emit("Login window timed out. Session saved as-is.")

                    self.done.emit(logged_in)

                finally:
                    page.close()
                    context.close()

        except Exception as e:
            self.failed.emit(str(e))


class Worker(QThread):
    log = Signal(str)
    row_update = Signal(int, str, str)
    progress = Signal(int)
    finished = Signal()
    failed = Signal(str)
    failures = Signal(list)   # list of still-failed items at the end of a run

    def __init__(self, bot_url, excel_path, output_root, headless,
                 skip_existing=True, fail_threshold=3, cooldown_seconds=300,
                 auto_retry=True, retry_items=None,
                 wait_round_sec=0, wait_question_sec=0, wait_supplier_sec=0,
                 rate=False, dry_run=True, openai_model="", openai_key=""):
        super().__init__()
        self.bot_url = bot_url
        self.excel_path = excel_path
        self.output_root = output_root
        self.headless = headless
        self.skip_existing = skip_existing
        self.fail_threshold = fail_threshold
        self.cooldown_seconds = cooldown_seconds
        self.auto_retry = auto_retry
        self.retry_items = retry_items      # None = full run; list = retry only
        self.wait_round_sec = wait_round_sec
        self.wait_question_sec = wait_question_sec
        self.wait_supplier_sec = wait_supplier_sec
        self.rate = rate
        self.dry_run = dry_run
        self.openai_model = openai_model
        self.openai_key = openai_key
        self.stop_requested = False
        self.consecutive_failures = 0
        self.failed_items = []
        self._answers_by_question = {}
        self._row = 0
        self._completed = 0
        self._total = 0
        # _resume is "set" while running, "cleared" while paused.
        self._resume = threading.Event()
        self._resume.set()

    def request_stop(self):
        self.stop_requested = True
        # Wake the thread if it's parked in a pause so it can exit.
        self._resume.set()

    def request_pause(self):
        self._resume.clear()

    def request_resume(self):
        self._resume.set()

    def _fmt_duration(self, seconds):
        seconds = int(seconds)
        if seconds >= 60:
            minutes, secs = divmod(seconds, 60)
            return f"{minutes}m {secs}s" if secs else f"{minutes}m"
        return f"{seconds}s"

    def _interruptible_sleep(self, seconds, label):
        """Wait, checking for Stop every second. Returns True if stopped."""
        if seconds <= 0:
            return False
        self.log.emit(f"Waiting {label}: {self._fmt_duration(seconds)}.")
        deadline = time.time() + seconds
        while time.time() < deadline:
            if self.stop_requested:
                return True
            time.sleep(1)
        return False

    def _pacing_wait(self, current, nxt):
        """Pause between two real bot interactions based on what changes next:
        a new supplier, a new question, or just the next round."""
        if nxt["supplier"] != current["supplier"]:
            seconds = self.wait_supplier_sec
            label = "between suppliers"
        elif nxt["question_index"] != current["question_index"]:
            seconds = self.wait_question_sec
            label = "between questions"
        else:
            seconds = self.wait_round_sec
            label = "after round"
        return self._interruptible_sleep(seconds, label)

    def _start_caffeinate(self):
        """Keep the Mac awake for the duration of the batch (no-op off macOS)."""
        if sys.platform != "darwin":
            return None
        try:
            proc = subprocess.Popen(["caffeinate", "-dims"])
            self.log.emit("Keeping the Mac awake for this batch.")
            return proc
        except Exception:
            return None

    def _stop_caffeinate(self, proc):
        if proc:
            try:
                proc.terminate()
            except Exception:
                pass

    def _wait_if_paused_or_stop(self):
        """Checkpoint between runs. Returns True if the worker should stop."""
        if self.stop_requested:
            return True
        if not self._resume.is_set():
            self.log.emit("Paused. Click Resume to continue.")
            self._resume.wait()
            if self.stop_requested:
                return True
            self.log.emit("Resumed.")
        return False

    def _register_failure(self):
        self.consecutive_failures += 1

    def _cooldown_if_needed(self):
        """If too many failures in a row, pause for the cooldown. Returns True if
        the user stopped during the wait."""
        if self.fail_threshold > 0 and self.consecutive_failures >= self.fail_threshold:
            minutes = self.cooldown_seconds / 60.0
            self.log.emit(
                f"{self.consecutive_failures} failures in a row — cooling down "
                f"{minutes:.0f} min to avoid throttling."
            )
            deadline = time.time() + self.cooldown_seconds
            next_note = time.time() + 30
            while time.time() < deadline:
                if self.stop_requested:
                    return True
                time.sleep(1)
                if time.time() >= next_note:
                    remaining = int(deadline - time.time())
                    if remaining > 0:
                        self.log.emit(f"Cooldown: ~{remaining}s remaining.")
                    next_note += 30
            self.consecutive_failures = 0
            self.log.emit("Cooldown complete. Resuming.")
        return False

    def _rate_answer(self, item, target_path, page):
        """Judge one downloaded answer with AI, write feedback into the Word file,
        and (unless dry-run) submit the rating on the site. Runs inside run_one
        while the page is still open."""
        if not docx_has_text(target_path):
            return  # empty answer; the empty-doc check will delete + retry

        key = (item["supplier"], item["question_index"])
        answer_text = docx_extract_text(target_path)
        siblings = self._answers_by_question.get(key, [])

        try:
            result = ai_judge(
                self.openai_key, self.openai_model, item["question"],
                item.get("approved", ""), item.get("expected_suppliers", ""),
                answer_text, siblings, target_supplier=item["supplier"]
            )
        except Exception as exc:
            self.log.emit(f"AI rating failed: {exc}")
            import traceback
            for ln in traceback.format_exc().strip().splitlines()[-5:]:
                self.log.emit(ln)
            return

        self._answers_by_question.setdefault(key, []).append(answer_text)
        try:
            append_qa_feedback(target_path, result)
        except Exception as exc:
            self.log.emit(f"Could not write feedback into the Word file: {exc}")

        missing = ", ".join(result["suppliers_missing"]) or "none"
        self.log.emit(
            f"AI verdict: {result['decision'].upper()} {result['score']}/10 "
            f"(missing suppliers: {missing})"
        )
        append_result_row(
            self.output_root, item["supplier"], item["question_index"],
            item["question"], item["round_number"],
            f"Rated {result['decision']} {result['score']}/10",
            result["feedback"][:250]
        )

        if not self.dry_run:
            try:
                shot = target_path.parent / f"{target_path.stem}_rating.png"
                submit_rating(page, result, self.log.emit, shot_path=str(shot))
            except Exception as exc:
                self.log.emit(f"Could not submit rating on the site: {exc}")

    def _attempt_round(self, runner, item, row):
        supplier = item["supplier"]
        question_index = item["question_index"]
        question = item["question"]
        round_number = item["round_number"]
        label = f'{supplier} | Q{question_index:02d} | Round {round_number:02d}'

        # Resume-skip, but only if the existing file actually has text.
        if self.skip_existing and self.retry_items is None:
            existing = existing_round_file(self.output_root, supplier, question_index, round_number)
            if existing and docx_has_text(existing):
                self.log.emit(f"Skipping (already downloaded): {existing}")
                self.row_update.emit(row, str(existing), "Skipped")
                append_result_row(self.output_root, supplier, question_index,
                                   question, round_number, "Skipped", existing)
                return "skipped"
            elif existing:
                self.log.emit(f"Existing file was empty, re-downloading: {existing}")
                try:
                    existing.unlink()
                except Exception:
                    pass

        self.log.emit(f"Starting: {label}")
        self.row_update.emit(row, label, "Running")

        last_error = ""
        for attempt in range(1, 4):
            if self.stop_requested:
                return "stopped"
            try:
                self.log.emit(f"Attempt {attempt}/3")
                rate_cb = None
                if self.rate:
                    rate_cb = lambda target, page: self._rate_answer(item, target, page)
                file_path = runner.run_one(
                    supplier, question_index, question, round_number,
                    rate_callback=rate_cb
                )

                # Reject empty/no-text downloads: delete and retry.
                if not docx_has_text(file_path):
                    try:
                        Path(file_path).unlink()
                    except Exception:
                        pass
                    last_error = "Downloaded Word file had no text (empty answer)."
                    self.log.emit(last_error)
                    self._register_failure()
                    if self._cooldown_if_needed():
                        return "stopped"
                    continue

                self.row_update.emit(row, str(file_path), "Downloaded")
                self.log.emit(f"Downloaded: {file_path}")
                append_result_row(self.output_root, supplier, question_index,
                                  question, round_number, "Downloaded", file_path)
                self.consecutive_failures = 0
                return "ok"

            except PlaywrightTimeoutError:
                last_error = "Timeout waiting for page, response, or download."
                self.log.emit(last_error)
                self._register_failure()
                if self._cooldown_if_needed():
                    return "stopped"

            except Exception as e:
                last_error = str(e)
                self.log.emit(f"Error: {last_error}")
                self._register_failure()
                if self._cooldown_if_needed():
                    return "stopped"

        self.row_update.emit(row, last_error, "Failed")
        append_result_row(self.output_root, supplier, question_index,
                          question, round_number, "Failed", last_error)
        self.failed_items.append({
            "supplier": supplier,
            "question_index": question_index,
            "question": question,
            "round_number": round_number,
        })
        return "failed"

    def _process_items(self, runner, items, count_progress):
        for idx, item in enumerate(items):
            if self._wait_if_paused_or_stop():
                return True
            status = self._attempt_round(runner, item, self._row)
            self._row += 1
            if status == "stopped":
                return True
            if count_progress and self._total:
                self._completed += 1
                self.progress.emit(int((self._completed / self._total) * 100))
            # Pace only after a real bot interaction (skipped rounds don't wait).
            if status in ("ok", "failed") and idx < len(items) - 1:
                if self._pacing_wait(item, items[idx + 1]):
                    return True
        return False

    def _finish(self, grouped_jobs, work_items):
        if self.retry_items is None:
            suppliers = list(grouped_jobs.keys())
        else:
            suppliers = sorted({it["supplier"] for it in work_items})

        for supplier in suppliers:
            if supplier not in grouped_jobs:
                continue
            if not supplier_has_downloads(self.output_root, supplier, grouped_jobs[supplier]):
                continue
            self.log.emit(f"Building combined report: {supplier}")
            try:
                report_path = build_combined_report_from_spec(
                    self.output_root, supplier, grouped_jobs[supplier]
                )
                self.log.emit(f"Combined report: {report_path}")
            except Exception as exc:
                self.log.emit(f"Could not build report for {supplier}: {exc}")

        self.failures.emit(list(self.failed_items))
        self.finished.emit()

    def run(self):
        caffeinate = self._start_caffeinate()
        try:
            all_jobs = load_questions(self.excel_path)
            grouped_jobs = group_jobs_by_supplier(all_jobs)

            # First-run browser setup (downloads Chromium once if needed).
            ensure_chromium(self.log.emit)

            runner = BotRunner(
                self.bot_url,
                self.output_root,
                self.headless,
                self.log.emit
            )

            if self.retry_items is not None:
                work_items = list(self.retry_items)
                self.log.emit(f"Retrying {len(work_items)} previously failed item(s).")
            else:
                work_items = []
                for job in all_jobs:
                    for round_number in range(1, job["repetitions"] + 1):
                        work_items.append({
                            "supplier": job["supplier"],
                            "question_index": job["question_index"],
                            "question": job["question"],
                            "round_number": round_number,
                            "approved": job.get("approved", ""),
                            "expected_suppliers": job.get("expected_suppliers", ""),
                        })

            self._total = len(work_items)
            self._completed = 0
            self._row = 0
            self.failed_items = []
            self.consecutive_failures = 0
            self._answers_by_question = {}

            stopped = self._process_items(runner, work_items, count_progress=True)
            if stopped:
                self.log.emit("Stopped by user.")
                self._finish(grouped_jobs, work_items)
                return

            # Automatic retry pass (full runs only).
            if self.auto_retry and self.retry_items is None and self.failed_items:
                retry_list = self.failed_items
                self.failed_items = []
                self.log.emit(f"Auto-retry pass for {len(retry_list)} failed item(s).")
                stopped = self._process_items(runner, retry_list, count_progress=False)
                if stopped:
                    self.log.emit("Stopped by user.")
                    self._finish(grouped_jobs, work_items)
                    return

            self._finish(grouped_jobs, work_items)

        except Exception as e:
            self.failed.emit(str(e))
        finally:
            self._stop_caffeinate(caffeinate)


class MainWindow(QMainWindow):
    def __init__(self):
        super().__init__()

        self.setWindowTitle("❄ Winter Bot Tester")
        self.resize(1400, 760)

        self.worker = None
        self.login_worker = None
        self.last_failed_items = []

        self.bot_url = QLineEdit()
        self.bot_url.setPlaceholderText("Paste bot URL here")

        self.excel_path = QLineEdit()
        self.output_folder = QLineEdit()

        self.headless = QCheckBox("Headless mode")
        self.headless.setChecked(False)

        self.skip_existing = QCheckBox("Skip rounds already downloaded today (resume)")
        self.skip_existing.setChecked(True)

        self.auto_retry = QCheckBox("Auto-retry failed items at the end")
        self.auto_retry.setChecked(True)

        self.fail_threshold_spin = QSpinBox()
        self.fail_threshold_spin.setRange(1, 50)
        self.fail_threshold_spin.setValue(3)

        self.cooldown_spin = QSpinBox()
        self.cooldown_spin.setRange(0, 180)
        self.cooldown_spin.setValue(5)

        self.wait_round_spin = QSpinBox()
        self.wait_round_spin.setRange(0, 9999)
        self.wait_round_spin.setValue(5)
        self.wait_round_unit = QComboBox()
        self.wait_round_unit.addItems(["seconds", "minutes"])

        self.wait_question_spin = QSpinBox()
        self.wait_question_spin.setRange(0, 9999)
        self.wait_question_spin.setValue(1)
        self.wait_question_unit = QComboBox()
        self.wait_question_unit.addItems(["seconds", "minutes"])
        self.wait_question_unit.setCurrentText("minutes")

        self.wait_supplier_spin = QSpinBox()
        self.wait_supplier_spin.setRange(0, 9999)
        self.wait_supplier_spin.setValue(5)
        self.wait_supplier_unit = QComboBox()
        self.wait_supplier_unit.addItems(["seconds", "minutes"])
        self.wait_supplier_unit.setCurrentText("minutes")

        self.rate_mode = QCheckBox("Rate responses with AI (training)")
        self.rate_mode.setChecked(False)
        self.dry_run = QCheckBox("Dry run (judge only, don't click thumbs/Save)")
        self.dry_run.setChecked(True)

        self.openai_model = QLineEdit()
        self.openai_model.setPlaceholderText("OpenAI model (e.g. gpt-4o-mini)")
        self.openai_key = QLineEdit()
        self.openai_key.setEchoMode(QLineEdit.Password)
        self.openai_key.setPlaceholderText("OpenAI API key (or set OPENAI_API_KEY)")

        self.login_btn = QPushButton("First-time login")
        self.start_btn = QPushButton("Start")
        self.pause_btn = QPushButton("Pause")
        self.pause_btn.setEnabled(False)
        self.stop_btn = QPushButton("Stop")
        self.stop_btn.setEnabled(False)
        self.retry_btn = QPushButton("Retry failed")
        self.retry_btn.setEnabled(False)

        self.progress = QProgressBar()

        self.table = QTableWidget(0, 3)
        self.table.setHorizontalHeaderLabels(["Run", "File / Error", "Status"])
        self.table.horizontalHeader().setSectionResizeMode(QHeaderView.Stretch)

        self.logs = QTextEdit()
        self.logs.setReadOnly(True)

        self.build_ui()
        self.connect_events()

        self.settings = QSettings("WinterBotTester", APP_NAME)
        self.load_settings()

    def build_ui(self):
        root = QWidget()
        main_layout = QHBoxLayout(root)

        left = QFrame()
        left.setFixedWidth(360)
        left_layout = QVBoxLayout(left)

        title = QLabel("❄ Winter Bot Tester")
        title.setStyleSheet("""
            font-size: 28px;
            font-weight: 700;
            color: #F9FAFB;
            padding-bottom: 5px;
        """)

        subtitle = QLabel("Supplier QA Automation Suite")
        subtitle.setStyleSheet("""
            font-size: 13px;
            color: #9CA3AF;
            padding-bottom: 15px;
        """)

        left_layout.addWidget(title)
        left_layout.addWidget(subtitle)

        left_layout.addWidget(QLabel("Bot URL"))
        left_layout.addWidget(self.bot_url)

        left_layout.addWidget(QLabel("Excel File"))
        excel_row = QHBoxLayout()
        excel_row.addWidget(self.excel_path)

        excel_btn = QPushButton("Browse")
        excel_btn.clicked.connect(self.pick_excel)
        excel_row.addWidget(excel_btn)

        left_layout.addLayout(excel_row)

        left_layout.addWidget(QLabel("Output Folder"))
        output_row = QHBoxLayout()
        output_row.addWidget(self.output_folder)

        output_btn = QPushButton("Browse")
        output_btn.clicked.connect(self.pick_output)
        output_row.addWidget(output_btn)

        left_layout.addLayout(output_row)

        left_layout.addWidget(self.headless)
        left_layout.addWidget(self.skip_existing)
        left_layout.addWidget(self.auto_retry)

        threshold_row = QHBoxLayout()
        threshold_row.addWidget(QLabel("Cooldown after N fails in a row:"))
        threshold_row.addWidget(self.fail_threshold_spin)
        left_layout.addLayout(threshold_row)

        cooldown_row = QHBoxLayout()
        cooldown_row.addWidget(QLabel("Cooldown length (minutes):"))
        cooldown_row.addWidget(self.cooldown_spin)
        left_layout.addLayout(cooldown_row)

        left_layout.addWidget(QLabel("Pacing (to avoid being flagged):"))

        wait_round_row = QHBoxLayout()
        wait_round_row.addWidget(QLabel("Wait after each round:"))
        wait_round_row.addWidget(self.wait_round_spin)
        wait_round_row.addWidget(self.wait_round_unit)
        left_layout.addLayout(wait_round_row)

        wait_question_row = QHBoxLayout()
        wait_question_row.addWidget(QLabel("Wait between questions:"))
        wait_question_row.addWidget(self.wait_question_spin)
        wait_question_row.addWidget(self.wait_question_unit)
        left_layout.addLayout(wait_question_row)

        wait_supplier_row = QHBoxLayout()
        wait_supplier_row.addWidget(QLabel("Wait between suppliers:"))
        wait_supplier_row.addWidget(self.wait_supplier_spin)
        wait_supplier_row.addWidget(self.wait_supplier_unit)
        left_layout.addLayout(wait_supplier_row)

        left_layout.addWidget(QLabel("AI rating (training):"))
        left_layout.addWidget(self.rate_mode)
        left_layout.addWidget(self.dry_run)
        left_layout.addWidget(self.openai_model)
        left_layout.addWidget(self.openai_key)

        left_layout.addWidget(self.login_btn)
        left_layout.addWidget(self.start_btn)
        left_layout.addWidget(self.pause_btn)
        left_layout.addWidget(self.stop_btn)
        left_layout.addWidget(self.retry_btn)
        left_layout.addStretch()

        center = QFrame()
        center_layout = QVBoxLayout(center)

        progress_title = QLabel("Execution Progress")
        progress_title.setStyleSheet("""
            font-size: 18px;
            font-weight: 600;
            color: #F9FAFB;
        """)

        center_layout.addWidget(progress_title)
        center_layout.addWidget(self.progress)
        center_layout.addWidget(self.table)

        right = QFrame()
        right.setFixedWidth(420)
        right_layout = QVBoxLayout(right)

        logs_title = QLabel("Live Logs")
        logs_title.setStyleSheet("""
            font-size: 18px;
            font-weight: 600;
            color: #F9FAFB;
        """)

        right_layout.addWidget(logs_title)
        right_layout.addWidget(self.logs)

        main_layout.addWidget(left)
        main_layout.addWidget(center)
        main_layout.addWidget(right)

        self.setCentralWidget(root)

        self.setStyleSheet("""
            QMainWindow {
                background: #111827;
            }

            QFrame {
                background: #1F2937;
                border-radius: 14px;
            }

            QLabel {
                color: #F9FAFB;
            }

            QLineEdit,
            QTextEdit,
            QTableWidget {
                background: #111827;
                color: #F9FAFB;
                border: 1px solid #374151;
                border-radius: 8px;
                padding: 8px;
            }

            QPushButton {
                background: #3B82F6;
                color: white;
                border: none;
                border-radius: 8px;
                padding: 10px;
                font-weight: 600;
            }

            QPushButton:hover {
                background: #60A5FA;
            }

            QPushButton:disabled {
                background: #4B5563;
            }

            QCheckBox {
                color: #F9FAFB;
            }

            QProgressBar {
                color: white;
                border: 1px solid #374151;
                border-radius: 8px;
                text-align: center;
                background: #111827;
                height: 24px;
            }

            QProgressBar::chunk {
                background: #60A5FA;
                border-radius: 8px;
            }

            QHeaderView::section {
                background: #374151;
                color: #F9FAFB;
                padding: 6px;
                border: none;
            }

            QSpinBox, QComboBox {
                background: #111827;
                color: #F9FAFB;
                border: 1px solid #374151;
                border-radius: 8px;
                padding: 4px 8px;
                min-height: 22px;
            }

            QSpinBox::up-button {
                subcontrol-origin: border;
                subcontrol-position: top right;
                width: 20px;
                border-left: 1px solid #374151;
            }

            QSpinBox::down-button {
                subcontrol-origin: border;
                subcontrol-position: bottom right;
                width: 20px;
                border-left: 1px solid #374151;
            }

            QSpinBox::up-button:hover,
            QSpinBox::down-button:hover {
                background: #374151;
            }

            QSpinBox::up-arrow {
                width: 0;
                height: 0;
                border-left: 4px solid transparent;
                border-right: 4px solid transparent;
                border-bottom: 6px solid #F9FAFB;
            }

            QSpinBox::down-arrow {
                width: 0;
                height: 0;
                border-left: 4px solid transparent;
                border-right: 4px solid transparent;
                border-top: 6px solid #F9FAFB;
            }

            QComboBox::drop-down {
                subcontrol-origin: border;
                subcontrol-position: center right;
                width: 22px;
                border-left: 1px solid #374151;
            }

            QComboBox::down-arrow {
                width: 0;
                height: 0;
                border-left: 4px solid transparent;
                border-right: 4px solid transparent;
                border-top: 6px solid #F9FAFB;
            }

            QComboBox QAbstractItemView {
                background: #1F2937;
                color: #F9FAFB;
                selection-background-color: #3B82F6;
                border: 1px solid #374151;
            }
        """)

    def connect_events(self):
        self.login_btn.clicked.connect(self.toggle_login)
        self.start_btn.clicked.connect(self.start)
        self.pause_btn.clicked.connect(self.toggle_pause)
        self.stop_btn.clicked.connect(self.stop)
        self.retry_btn.clicked.connect(self.retry_failed)

    def load_settings(self):
        self.bot_url.setText(self.settings.value("bot_url", "", type=str))
        self.excel_path.setText(self.settings.value("excel_path", "", type=str))
        self.output_folder.setText(self.settings.value("output_folder", "", type=str))
        self.headless.setChecked(self.settings.value("headless", False, type=bool))
        self.skip_existing.setChecked(self.settings.value("skip_existing", True, type=bool))
        self.auto_retry.setChecked(self.settings.value("auto_retry", True, type=bool))
        self.fail_threshold_spin.setValue(self.settings.value("fail_threshold", 3, type=int))
        self.cooldown_spin.setValue(self.settings.value("cooldown_minutes", 5, type=int))
        self.wait_round_spin.setValue(self.settings.value("wait_round_value", 5, type=int))
        self.wait_round_unit.setCurrentText(self.settings.value("wait_round_unit", "seconds", type=str))
        self.wait_question_spin.setValue(self.settings.value("wait_question_value", 1, type=int))
        self.wait_question_unit.setCurrentText(self.settings.value("wait_question_unit", "minutes", type=str))
        self.wait_supplier_spin.setValue(self.settings.value("wait_supplier_value", 5, type=int))
        self.wait_supplier_unit.setCurrentText(self.settings.value("wait_supplier_unit", "minutes", type=str))
        self.rate_mode.setChecked(self.settings.value("rate_mode", False, type=bool))
        self.dry_run.setChecked(self.settings.value("dry_run", True, type=bool))
        self.openai_model.setText(self.settings.value("openai_model", "gpt-4o-mini", type=str))
        self.openai_key.setText(self.settings.value("openai_key", "", type=str))

    def save_settings(self):
        self.settings.setValue("bot_url", self.bot_url.text().strip())
        self.settings.setValue("excel_path", self.excel_path.text().strip())
        self.settings.setValue("output_folder", self.output_folder.text().strip())
        self.settings.setValue("headless", self.headless.isChecked())
        self.settings.setValue("skip_existing", self.skip_existing.isChecked())
        self.settings.setValue("auto_retry", self.auto_retry.isChecked())
        self.settings.setValue("fail_threshold", self.fail_threshold_spin.value())
        self.settings.setValue("cooldown_minutes", self.cooldown_spin.value())
        self.settings.setValue("wait_round_value", self.wait_round_spin.value())
        self.settings.setValue("wait_round_unit", self.wait_round_unit.currentText())
        self.settings.setValue("wait_question_value", self.wait_question_spin.value())
        self.settings.setValue("wait_question_unit", self.wait_question_unit.currentText())
        self.settings.setValue("wait_supplier_value", self.wait_supplier_spin.value())
        self.settings.setValue("wait_supplier_unit", self.wait_supplier_unit.currentText())
        self.settings.setValue("rate_mode", self.rate_mode.isChecked())
        self.settings.setValue("dry_run", self.dry_run.isChecked())
        self.settings.setValue("openai_model", self.openai_model.text().strip())
        self.settings.setValue("openai_key", self.openai_key.text())

    def closeEvent(self, event):
        self.save_settings()
        super().closeEvent(event)

    def pick_excel(self):
        path, _ = QFileDialog.getOpenFileName(
            self,
            "Select Excel File",
            "",
            "Excel Files (*.xlsx *.xls)"
        )

        if path:
            self.excel_path.setText(path)

    def pick_output(self):
        path = QFileDialog.getExistingDirectory(self, "Select Output Folder")

        if path:
            self.output_folder.setText(path)

    def start(self):
        bot_url = self.bot_url.text().strip()
        excel_path = self.excel_path.text().strip()
        output_folder = self.output_folder.text().strip()

        if not bot_url or not excel_path or not output_folder:
            QMessageBox.warning(
                self,
                "Missing data",
                "Please add bot URL, Excel file, and output folder."
            )
            return

        try:
            jobs = load_questions(excel_path)
            total = sum(job["repetitions"] for job in jobs)
            self.table.setRowCount(total)
        except Exception as e:
            QMessageBox.critical(self, "Excel error", str(e))
            return

        if self.rate_mode.isChecked():
            key = self._openai_key()
            looks_wrong = key and (
                len(key) > 300
                or any(c.isspace() for c in key)
                or any(ord(c) > 127 for c in key)
            )
            if looks_wrong:
                QMessageBox.warning(
                    self, "AI rating",
                    f"The OpenAI API key field doesn't look like a key.\n\n"
                    f"It currently holds {len(key)} characters, with spaces or "
                    f"non-standard characters in it — that looks like pasted text, "
                    f"not a key.\n\nClear that field and paste ONLY your OpenAI key "
                    f"(it starts with 'sk-'), or leave it empty to use plain "
                    f"supplier-surfacing checks without OpenAI."
                )
                return

        self.logs.clear()
        self.add_log(f"{APP_NAME} v{APP_VERSION}")
        self.progress.setValue(0)
        self.save_settings()
        self.last_failed_items = []

        if self.rate_mode.isChecked():
            mode = "DRY RUN (no clicks)" if self.dry_run.isChecked() else "LIVE (will submit)"
            self.add_log(f"AI rating is ON — {mode}.")
            keylen = len(self._openai_key())
            if keylen:
                self.add_log(f"Supplier-surfacing check; OpenAI feedback ON ({keylen}-char key).")
            else:
                self.add_log("Supplier-surfacing check; no OpenAI key (plain feedback).")

        self._launch_worker(retry_items=None)

    def retry_failed(self):
        if not self.last_failed_items:
            return

        if not (self.bot_url.text().strip()
                and self.excel_path.text().strip()
                and self.output_folder.text().strip()):
            QMessageBox.warning(self, "Missing data",
                                "Please add bot URL, Excel file, and output folder.")
            return

        items = list(self.last_failed_items)
        self.add_log(f"Retrying {len(items)} failed item(s)...")
        self.progress.setValue(0)
        if self.table.rowCount() < len(items):
            self.table.setRowCount(len(items))

        self._launch_worker(retry_items=items)

    def _launch_worker(self, retry_items=None):
        self.worker = Worker(
            self.bot_url.text().strip(),
            self.excel_path.text().strip(),
            self.output_folder.text().strip(),
            self.headless.isChecked(),
            self.skip_existing.isChecked(),
            self.fail_threshold_spin.value(),
            self.cooldown_spin.value() * 60,
            self.auto_retry.isChecked(),
            retry_items,
            self._wait_seconds(self.wait_round_spin, self.wait_round_unit),
            self._wait_seconds(self.wait_question_spin, self.wait_question_unit),
            self._wait_seconds(self.wait_supplier_spin, self.wait_supplier_unit),
            self.rate_mode.isChecked(),
            self.dry_run.isChecked(),
            self.openai_model.text().strip(),
            self._openai_key()
        )

        self.worker.log.connect(self.add_log)
        self.worker.row_update.connect(self.update_row)
        self.worker.progress.connect(self.progress.setValue)
        self.worker.failures.connect(self.store_failures)
        self.worker.finished.connect(self.finished)
        self.worker.failed.connect(self.failed)

        self.login_btn.setEnabled(False)
        self.start_btn.setEnabled(False)
        self.retry_btn.setEnabled(False)
        self.pause_btn.setEnabled(True)
        self.pause_btn.setText("Pause")
        self.stop_btn.setEnabled(True)

        self.worker.start()

    def store_failures(self, items):
        self.last_failed_items = items

    def _wait_seconds(self, spin, unit_combo):
        value = spin.value()
        return value * 60 if unit_combo.currentText() == "minutes" else value

    def _openai_key(self):
        return self.openai_key.text().strip() or os.environ.get("OPENAI_API_KEY", "")

    def toggle_login(self):
        # If a login session is already running, finish and save it.
        if self.login_worker and self.login_worker.isRunning():
            self.login_worker.request_finish()
            self.add_log("Finishing login...")
            return

        bot_url = self.bot_url.text().strip()
        if not bot_url:
            QMessageBox.warning(self, "Missing data", "Please paste the bot URL first.")
            return

        self.logs.clear()
        self.add_log("Starting first-time login.")

        self.login_worker = LoginWorker(bot_url)
        self.login_worker.log.connect(self.add_log)
        self.login_worker.done.connect(self.login_done)
        self.login_worker.failed.connect(self.login_failed)

        self.start_btn.setEnabled(False)
        self.pause_btn.setEnabled(False)
        self.stop_btn.setEnabled(False)
        self.login_btn.setText("Finish login & save")

        self.login_worker.start()

    def login_done(self, detected):
        self.login_btn.setText("First-time login")
        self.start_btn.setEnabled(True)
        if detected:
            self.add_log("Login saved. You can run batches now (headless is fine).")
        else:
            self.add_log("Login session closed.")

    def login_failed(self, error):
        self.login_btn.setText("First-time login")
        self.start_btn.setEnabled(True)
        QMessageBox.critical(self, "Login error", error)
        self.add_log(error)

    def toggle_pause(self):
        if not self.worker:
            return
        if self.pause_btn.text() == "Pause":
            self.worker.request_pause()
            self.pause_btn.setText("Resume")
            self.add_log("Pause requested (takes effect after the current run finishes).")
        else:
            self.worker.request_resume()
            self.pause_btn.setText("Pause")
            self.add_log("Resume requested.")

    def stop(self):
        if self.worker:
            self.worker.request_stop()
            self.add_log("Stopping after current action finishes...")

    def add_log(self, message):
        timestamp = datetime.now().strftime("%H:%M:%S")
        self.logs.append(f"[{timestamp}] {message}")

    def update_row(self, row, detail, status):
        if row >= self.table.rowCount():
            self.table.setRowCount(row + 1)
        self.table.setItem(row, 0, QTableWidgetItem(str(row + 1)))
        self.table.setItem(row, 1, QTableWidgetItem(detail))
        self.table.setItem(row, 2, QTableWidgetItem(status))

    def finished(self):
        self.login_btn.setEnabled(True)
        self.start_btn.setEnabled(True)
        self.pause_btn.setEnabled(False)
        self.pause_btn.setText("Pause")
        self.stop_btn.setEnabled(False)
        self.retry_btn.setEnabled(bool(self.last_failed_items))
        self.add_log("Process completed.")
        if self.last_failed_items:
            self.add_log(
                f"{len(self.last_failed_items)} item(s) still failed — "
                "'Retry failed' is available."
            )

    def failed(self, error):
        self.login_btn.setEnabled(True)
        self.start_btn.setEnabled(True)
        self.pause_btn.setEnabled(False)
        self.pause_btn.setText("Pause")
        self.stop_btn.setEnabled(False)
        self.retry_btn.setEnabled(bool(self.last_failed_items))
        QMessageBox.critical(self, "Error", error)
        self.add_log(error)


if __name__ == "__main__":
    app = QApplication(sys.argv)
    app.setApplicationName(APP_NAME)

    window = MainWindow()
    window.show()

    sys.exit(app.exec())